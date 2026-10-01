import os
from docx import Document
from loguru import logger


def parse_resume(file_path: str) -> str:
    """Extract all text from a .docx or .pdf resume file."""
    ext = os.path.splitext(file_path)[1].lower()

    if ext == ".docx":
        return _parse_docx(file_path)
    elif ext == ".pdf":
        return _parse_pdf(file_path)
    else:
        logger.error(f"Unsupported file format: {ext}")
        return ""


def _parse_docx(file_path: str) -> str:
    """Extract text from a .docx file."""
    try:
        doc = Document(file_path)
        full_text = []

        for paragraph in doc.paragraphs:
            if paragraph.text.strip():
                full_text.append(paragraph.text.strip())

        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        full_text.append(cell.text.strip())

        resume_text = "\n".join(full_text)
        logger.info(f"Parsed DOCX resume: {len(resume_text)} characters from {file_path}")
        return resume_text

    except Exception as e:
        logger.error(f"Failed to parse DOCX: {e}")
        return ""


def _parse_pdf(file_path: str) -> str:
    """Extract text from a PDF file."""
    try:
        import fitz  # pymupdf

        doc = fitz.open(file_path)
        full_text = []

        for page in doc:
            text = page.get_text()
            if text.strip():
                full_text.append(text.strip())

        doc.close()

        resume_text = "\n".join(full_text)
        logger.info(f"Parsed PDF resume: {len(resume_text)} characters from {file_path}")
        return resume_text

    except Exception as e:
        logger.error(f"Failed to parse PDF: {e}")
        return ""


if __name__ == "__main__":
    import sys
    if len(sys.argv) > 1:
        text = parse_resume(sys.argv[1])
        print(text)
    else:
        print("Usage: python -m jobhunter.resume.parser <path_to_resume.docx or .pdf>")